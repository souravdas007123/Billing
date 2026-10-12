"""Reading a list of parties from a CSV / Excel file or from rows pasted out of Excel.

Pure functions (no Django), so they are easy to test. The admin page in admin.py does the saving.
"""
import csv
import io
import re

from .calc import is_valid_gstin

MAX_ROWS = 5000
MAX_BYTES = 3 * 1024 * 1024
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

FIELDS = ["name", "address", "gstin", "state", "phone", "email"]
# Header names people really use, compared after lower-casing and dropping spaces and punctuation.
ALIASES = {
    "name": ["name", "party", "partyname", "customer", "customername", "company", "companyname", "buyer", "buyername", "ledger", "ledgername", "firmname"],
    "address": ["address", "addr", "partyaddress", "billingaddress", "customeraddress"],
    "gstin": ["gstin", "gstno", "gstnumber", "gstinno", "gstinnumber", "gstinuin", "gst", "gstid", "partygstin"],
    "state": ["state", "statename", "partystate", "placeofsupply"],
    "phone": ["phone", "phoneno", "phonenumber", "mobile", "mobileno", "mobilenumber", "contact", "contactno", "telephone", "whatsapp", "cell"],
    "email": ["email", "emailid", "emailaddress", "mail", "eemail"],
}
SAMPLE_CSV = (
    "name,address,gstin,state,phone,email\n"
    '"Alpha Traders","12 Park Street, Kolkata 700016","19AAAAA0000A1Z5","West Bengal","9830000000","accounts@alpha.example"\n'
    '"Beta Stores (no GSTIN)","5 Market Road, Pune","","Maharashtra","9890000000",""\n'
)


class ImportFileError(Exception):
    """The file as a whole cannot be used (wrong type, too big, no party-name column ...)."""


def _key(text):
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _cell(value):
    """A spreadsheet cell as clean text (9735752367.0 -> '9735752367')."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---------- reading the file ----------
def read_csv_text(text):
    text = text.lstrip("﻿")
    sample = text[:4096]
    delimiter = max((",", "\t", ";", "|"), key=sample.count)  # the separator that appears most
    return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportFileError("Reading Excel files needs the openpyxl package (pip install openpyxl). Or save the sheet as CSV.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("This does not look like a valid .xlsx file.")
    ws = wb.worksheets[0]
    return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]


def read_table(filename, data):
    """Rows of text from an uploaded file (.csv, .txt, .tsv, .xlsx)."""
    if len(data) > MAX_BYTES:
        raise ImportFileError("The file is too big (more than 3 MB).")
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return read_xlsx(data)
    if name.endswith(".xls"):
        raise ImportFileError("Old .xls files are not supported. Save it as .xlsx or CSV and upload again.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")  # Excel's "CSV" on Windows
    return read_csv_text(text)


# ---------- turning rows into parties ----------
def _find_header(table):
    """The header row (a title row above it is skipped) and the column index of each field."""
    for i, row in enumerate(table[:10]):
        cols = {}
        for j, head in enumerate(row):
            k = _key(head)
            for field, names in ALIASES.items():
                if k in names and field not in cols:
                    cols[field] = j
        if "name" in cols:
            return i, cols
    raise ImportFileError(
        "Could not find a column for the party name. The first row must have headings such as: "
        "name, address, gstin, state, phone, email."
    )


def _state(value, gstin, states):
    """(state name, warning). `states` maps a state name to its 2-digit code."""
    by_lower = {n.lower(): n for n in states}
    by_code = {c: n for n, c in states.items()}
    v = (value or "").strip()
    if v:
        if v.lower() in by_lower:
            return by_lower[v.lower()], ""
        if v.isdigit() and v.zfill(2) in by_code:
            return by_code[v.zfill(2)], ""
    name = by_code.get(gstin[:2], "") if len(gstin) >= 2 else ""
    if v:
        return name, f'state "{v}" is not recognised' + (" (taken from the GSTIN)" if name else " (left blank)")
    return name, ""


def build_parties(table, states):
    """Returns (parties, file_notes). Every party is a dict with the six fields plus row, warnings and error."""
    if not table:
        raise ImportFileError("The file is empty.")
    head, cols = _find_header(table)
    body = table[head + 1:]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
    parties, seen = [], {}
    for offset, row in enumerate(body):
        if not any(c.strip() for c in row):
            continue
        get = lambda f: row[cols[f]].strip() if f in cols and cols[f] < len(row) else ""
        p = {"row": head + offset + 2, "warnings": [], "error": ""}
        p["name"] = re.sub(r"\s+", " ", get("name"))
        p["address"] = get("address")
        gstin = re.sub(r"\s", "", get("gstin")).upper()
        p["gstin"] = gstin
        p["phone"] = get("phone")
        email = get("email")
        p["email"] = email
        if not p["name"]:
            p["error"] = "no party name"
        elif len(p["name"]) > 200:
            p["error"] = "party name is longer than 200 characters"
        elif p["name"].lower() in seen:
            p["error"] = f"same name as row {seen[p['name'].lower()]} of this file"
        else:
            seen[p["name"].lower()] = p["row"]
        if gstin and (len(gstin) != 15 or not is_valid_gstin(gstin)):
            p["warnings"].append("GSTIN looks invalid (saved as typed)")
        p["state"], warn = _state(get("state"), gstin, states)
        if warn:
            p["warnings"].append(warn)
        if len(p["phone"]) > 40:
            p["phone"] = p["phone"][:40]
            p["warnings"].append("phone shortened to 40 characters")
        if email and (len(email) > 254 or not EMAIL_RE.match(email)):
            p["warnings"].append(f'email "{email}" is not valid (left blank)')
            p["email"] = ""
        parties.append(p)
    if not parties:
        raise ImportFileError("No party rows were found below the heading row.")
    return parties, []
